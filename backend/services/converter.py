"""
보고서 포맷 변환 서비스
Markdown → PDF (WeasyPrint) / DOCX (python-docx)
"""
import io
import logging
import re

import markdown as md_lib
from weasyprint import HTML
from docx import Document
from docx.shared import Pt, RGBColor
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

logger = logging.getLogger(__name__)

# ── PDF CSS ────────────────────────────────────────────────────────────────────
_PDF_CSS = """
@font-face {
    font-family: 'NanumGothic';
    src: local('NanumGothic'), local('Nanum Gothic');
}

* { box-sizing: border-box; }

body {
    font-family: 'NanumGothic', 'Nanum Gothic', 'DejaVu Sans', sans-serif;
    font-size: 11pt;
    line-height: 1.9;
    color: #1e293b;
    margin: 0;
    padding: 0;
}

@page {
    margin: 25mm 20mm 25mm 20mm;
    @bottom-center {
        content: counter(page) " / " counter(pages);
        font-size: 9pt;
        color: #94a3b8;
    }
}

h1 {
    font-size: 18pt;
    font-weight: bold;
    color: #0f172a;
    margin-top: 0;
    margin-bottom: 16px;
    padding-bottom: 8px;
    border-bottom: 2px solid #3b82f6;
}

h2 {
    font-size: 14pt;
    font-weight: bold;
    color: #1e40af;
    margin-top: 28px;
    margin-bottom: 10px;
    padding-bottom: 4px;
    border-bottom: 1px solid #e2e8f0;
}

h3 {
    font-size: 12pt;
    font-weight: bold;
    color: #1e293b;
    margin-top: 20px;
    margin-bottom: 6px;
}

p {
    margin: 0 0 12px 0;
}

ul, ol {
    margin: 0 0 12px 0;
    padding-left: 24px;
}

li {
    margin-bottom: 4px;
}

strong { font-weight: bold; }
em     { font-style: italic; }

code {
    font-family: monospace;
    background: #f1f5f9;
    padding: 1px 4px;
    border-radius: 3px;
    font-size: 10pt;
}

pre {
    background: #f8fafc;
    border: 1px solid #e2e8f0;
    border-radius: 6px;
    padding: 12px;
    overflow-x: auto;
    font-size: 9.5pt;
    margin-bottom: 12px;
}

blockquote {
    border-left: 3px solid #3b82f6;
    margin: 0 0 12px 0;
    padding: 4px 16px;
    color: #475569;
    background: #f8fafc;
}

table {
    border-collapse: collapse;
    width: 100%;
    margin-bottom: 12px;
    font-size: 10pt;
}

th {
    background: #1e40af;
    color: white;
    padding: 8px 12px;
    text-align: left;
}

td {
    padding: 6px 12px;
    border: 1px solid #e2e8f0;
}

tr:nth-child(even) td {
    background: #f8fafc;
}
"""


def md_to_pdf(md_text: str) -> bytes:
    """Markdown 텍스트 → PDF 바이트"""
    try:
        # extra 안에 tables·fenced_code 포함 → 중복 지정하면 오류 날 수 있어서 분리
        html_body = md_lib.markdown(
            md_text,
            extensions=["tables", "fenced_code", "nl2br"],
        )
    except Exception as e:
        logger.error("[converter] markdown 파싱 실패: %s", e)
        # 파싱 실패 시 plain text 그대로 감싸서 진행
        html_body = f"<pre>{md_text}</pre>"

    full_html = (
        '<!DOCTYPE html>\n'
        '<html lang="ko">\n'
        '<head>\n'
        '  <meta charset="utf-8">\n'
        f'  <style>{_PDF_CSS}</style>\n'
        '</head>\n'
        f'<body>{html_body}</body>\n'
        '</html>'
    )

    try:
        return HTML(string=full_html).write_pdf()
    except Exception as e:
        logger.error("[converter] WeasyPrint PDF 변환 실패: %s", e, exc_info=True)
        raise


# ── DOCX 변환 ──────────────────────────────────────────────────────────────────

def _set_cell_bg(cell, hex_color: str):
    """표 셀 배경색 설정 헬퍼"""
    tc   = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd  = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    shd.set(qn("w:val"),  "clear")
    tcPr.append(shd)


def md_to_docx(md_text: str) -> bytes:
    """Markdown 텍스트 → DOCX 바이트"""
    doc = Document()

    # ── 기본 글꼴 설정 ────────────────────────────────────────────────────────
    normal = doc.styles["Normal"]
    normal.font.name = "Nanum Gothic"
    normal.font.size = Pt(11)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Nanum Gothic")

    # ── Markdown 파싱 ─────────────────────────────────────────────────────────
    lines   = md_text.splitlines()
    i       = 0
    in_list = False  # 연속 리스트 추적용

    while i < len(lines):
        raw  = lines[i]
        line = raw.rstrip()

        # 빈 줄
        if not line.strip():
            in_list = False
            i += 1
            continue

        # 제목
        if line.startswith("#### "):
            p = doc.add_heading(line[5:].strip(), level=4)
        elif line.startswith("### "):
            p = doc.add_heading(line[4:].strip(), level=3)
        elif line.startswith("## "):
            p = doc.add_heading(line[3:].strip(), level=2)
        elif line.startswith("# "):
            p = doc.add_heading(line[2:].strip(), level=1)

        # 수평선
        elif re.fullmatch(r"[-*_]{3,}", line.strip()):
            p = doc.add_paragraph()
            p.add_run("─" * 50)

        # 불릿 리스트
        elif re.match(r"^[-*+] ", line):
            doc.add_paragraph(line[2:].strip(), style="List Bullet")
            in_list = True

        # 번호 리스트
        elif re.match(r"^\d+\. ", line):
            text = re.sub(r"^\d+\. ", "", line)
            doc.add_paragraph(text.strip(), style="List Number")
            in_list = True

        # 인용
        elif line.startswith("> "):
            p = doc.add_paragraph(line[2:].strip())
            pf = p.paragraph_format
            pf.left_indent = Pt(18)
            run = p.runs[0] if p.runs else p.add_run()
            run.font.color.rgb = RGBColor(0x47, 0x55, 0x69)

        # 일반 문단 (인라인 볼드/이탤릭 처리)
        else:
            p = doc.add_paragraph()
            _add_inline(p, line.strip())

        i += 1

    bio = io.BytesIO()
    doc.save(bio)
    bio.seek(0)
    return bio.read()


# 인라인 마크다운(볼드·이탤릭) 처리 헬퍼
_INLINE_RE = re.compile(
    r"(\*\*\*(?P<bi>.+?)\*\*\*)"   # ***볼드이탤릭***
    r"|(\*\*(?P<b>.+?)\*\*)"       # **볼드**
    r"|(\*(?P<i>.+?)\*)"           # *이탤릭*
    r"|(`(?P<c>.+?)`)"             # `코드`
)


def _add_inline(paragraph, text: str):
    pos = 0
    for m in _INLINE_RE.finditer(text):
        if m.start() > pos:
            paragraph.add_run(text[pos:m.start()])

        run = paragraph.add_run(
            m.group("bi") or m.group("b") or m.group("i") or m.group("c")
        )
        if m.group("bi"):
            run.bold = True
            run.italic = True
        elif m.group("b"):
            run.bold = True
        elif m.group("i"):
            run.italic = True
        elif m.group("c"):
            run.font.name = "Courier New"
            run.font.size = Pt(10)

        pos = m.end()

    if pos < len(text):
        paragraph.add_run(text[pos:])
