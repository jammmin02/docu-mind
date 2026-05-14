"""
파일 형식별 텍스트 추출 서비스

PDF 처리 전략:
  1) PyMuPDF로 페이지별 텍스트 추출
  2) pdfplumber로 동일 페이지의 표 추출 → Markdown 변환
     - 표 영역(bbox) 내 텍스트는 일반 텍스트에서 제거 (중복 방지)
  3) 추출 텍스트가 OCR_THRESHOLD(50자) 미만인 페이지 → pytesseract OCR fallback
  4) 각 페이지를 PageResult(content, metadata)로 반환

비-PDF 파일은 기존과 동일하게 단일 또는 논리 단위별 텍스트 반환.
"""
from __future__ import annotations

import io
import logging
import re
from dataclasses import dataclass, field
from typing import Any

logger = logging.getLogger(__name__)

# MuPDF C 라이브러리의 stderr 직접 출력 경고 전역 억제
# 손상된 PDF에서 발생하는 경고가 애플리케이션 로그를 오염시키는 것을 줄이기 위함
try:
    import fitz as _fitz

    _fitz.TOOLS.mupdf_display_errors(False)
except Exception:
    pass

# OCR fallback 기준: 페이지에서 추출된 텍스트가 이 글자 수 미만이면 OCR 시도
OCR_THRESHOLD = 50


# ── 반환 타입 ──────────────────────────────────────────────────────────────────

@dataclass
class PageResult:
    """한 페이지 또는 논리적 블록의 파싱 결과."""

    content: str
    metadata: dict[str, Any] = field(default_factory=dict)
    # metadata 키:
    #   page_number : int | None — 1-based 페이지 번호
    #   has_table   : bool       — 표 포함 여부
    #   section     : str | None — 감지된 섹션 제목
    #   ocr_applied : bool       — OCR fallback 적용 여부


# ── 공개 API ──────────────────────────────────────────────────────────────────

def parse_file(contents: bytes, file_type: str) -> list[PageResult]:
    """
    파일을 파싱하여 PageResult 리스트로 반환.
    PDF/PPTX/XLSX는 페이지 또는 시트 단위, 나머지는 단일 PageResult로 반환한다.
    """
    normalized_type = file_type.lower().lstrip(".").strip()

    parsers = {
        "pdf": _parse_pdf,
        "docx": _parse_docx,
        "txt": _parse_txt,
        "xlsx": _parse_xlsx,
        "csv": _parse_csv,
        "pptx": _parse_pptx,
    }

    parser = parsers.get(normalized_type)
    if parser is None:
        raise ValueError(f"지원하지 않는 파일 형식입니다: {file_type}")

    pages = parser(contents)

    for page in pages:
        page.content = _strip_nul(page.content)
        page.metadata = {
            k: _strip_nul(v) if isinstance(v, str) else v
            for k, v in page.metadata.items()
        }

    total_text = "".join(p.content for p in pages).strip()
    if not total_text:
        raise ValueError("파일에서 텍스트를 추출할 수 없습니다")

    return pages


def _strip_nul(text: str) -> str:
    """PostgreSQL text 필드에 저장할 수 없는 NUL 바이트 제거."""
    return text.replace("\x00", "") if text else text


def pages_to_text(pages: list[PageResult]) -> str:
    """PageResult 리스트를 단일 텍스트로 합산한다."""
    return "\n\n".join(p.content for p in pages if p.content.strip())


# ── PDF 파싱 ──────────────────────────────────────────────────────────────────

def _is_slide_pdf(mupdf_doc) -> bool:
    """
    PDF가 슬라이드 형식인지 추정한다.
    가로 > 세로인 페이지가 전체의 70% 이상이면 슬라이드로 간주한다.
    """
    if len(mupdf_doc) == 0:
        return False

    landscape_count = sum(
        1
        for i in range(len(mupdf_doc))
        if mupdf_doc[i].rect.width > mupdf_doc[i].rect.height
    )

    return landscape_count / len(mupdf_doc) >= 0.7


def _parse_pdf(contents: bytes) -> list[PageResult]:
    import fitz
    import pdfplumber

    results: list[PageResult] = []
    mupdf_doc = None

    try:
        mupdf_doc = fitz.open(stream=contents, filetype="pdf")
        pdf_bytes = io.BytesIO(contents)

        is_slide = _is_slide_pdf(mupdf_doc)
        ocr_threshold = 20 if is_slide else OCR_THRESHOLD

        if is_slide:
            logger.info("[parser] detected slide-format PDF, OCR threshold=%d", ocr_threshold)

        with pdfplumber.open(pdf_bytes) as plumber_doc:
            page_count = min(len(mupdf_doc), len(plumber_doc.pages))

            for page_idx in range(page_count):
                mupdf_page = mupdf_doc[page_idx]
                plumber_page = plumber_doc.pages[page_idx]
                page_number = page_idx + 1

                tables, table_bboxes = _extract_tables(plumber_page)
                plain_text = _extract_text_excluding_tables(mupdf_page, table_bboxes)

                ocr_applied = False
                if len(plain_text.strip()) < ocr_threshold and not tables:
                    ocr_text = _ocr_page(mupdf_page)
                    if ocr_text.strip():
                        plain_text = ocr_text
                        ocr_applied = True
                        logger.info("[parser] OCR applied page=%d", page_number)

                section = _detect_section(mupdf_page, plain_text)

                parts: list[str] = []
                if plain_text.strip():
                    parts.append(plain_text.strip())
                parts.extend(tables)

                content = "\n\n".join(parts)
                if not content.strip():
                    continue

                results.append(
                    PageResult(
                        content=content,
                        metadata={
                            "page_number": page_number,
                            "has_table": len(tables) > 0,
                            "section": section,
                            "ocr_applied": ocr_applied,
                        },
                    )
                )

    finally:
        if mupdf_doc is not None:
            mupdf_doc.close()

    return results


# ── PDF 헬퍼 함수들 ───────────────────────────────────────────────────────────

def _extract_tables(plumber_page) -> tuple[list[str], list[tuple[float, float, float, float]]]:
    """
    pdfplumber로 표를 추출하고 Markdown 문자열 및 bbox를 반환한다.
    bbox는 pdfplumber 기준 좌표지만 일반적인 PDF 좌표계에서는 PyMuPDF block 좌표와 비교 가능하다.
    단, PDF에 따라 미세한 오차가 있을 수 있으므로 overlap 판정에 여유값을 둔다.
    """
    tables_md: list[str] = []
    table_bboxes: list[tuple[float, float, float, float]] = []

    try:
        for table in plumber_page.extract_tables() or []:
            if not table:
                continue

            md = _table_to_markdown(table)
            if md:
                tables_md.append(md)

        for tbl_obj in plumber_page.find_tables() or []:
            if getattr(tbl_obj, "bbox", None):
                table_bboxes.append(tbl_obj.bbox)

    except Exception as e:
        logger.warning("[parser] 표 추출 실패: %s", e)

    return tables_md, table_bboxes


def _table_to_markdown(table: list[list[Any]]) -> str:
    """2D 리스트를 Markdown 표 문자열로 변환한다."""
    if not table:
        return ""

    rows = [
        [str(cell).strip() if cell is not None else "" for cell in row]
        for row in table
    ]
    rows = [row for row in rows if any(cell for cell in row)]

    if not rows:
        return ""

    col_count = max(len(row) for row in rows)
    rows = [row + [""] * (col_count - len(row)) for row in rows]

    header = rows[0]
    body = rows[1:]
    sep = ["---"] * col_count

    lines = [
        "| " + " | ".join(header) + " |",
        "| " + " | ".join(sep) + " |",
    ]

    for row in body:
        lines.append("| " + " | ".join(row) + " |")

    return "\n".join(lines)


def _extract_text_excluding_tables(
    mupdf_page,
    table_bboxes: list[tuple[float, float, float, float]],
) -> str:
    """PyMuPDF 페이지에서 텍스트를 추출하되 표 영역과 겹치는 블록은 제외한다."""
    if not table_bboxes:
        return mupdf_page.get_text()

    blocks = mupdf_page.get_text("blocks")
    kept_parts: list[str] = []
    margin = 2.0

    for block in blocks:
        if len(block) < 5:
            continue

        x0, y0, x1, y1, text = block[:5]
        if not str(text).strip():
            continue

        in_table = False
        for tx0, ttop, tx1, tbottom in table_bboxes:
            overlap_x = x0 < tx1 + margin and x1 > tx0 - margin
            overlap_y = y0 < tbottom + margin and y1 > ttop - margin

            if overlap_x and overlap_y:
                in_table = True
                break

        if not in_table:
            kept_parts.append(str(text))

    return "\n".join(kept_parts)


def _ocr_page(mupdf_page) -> str:
    """PyMuPDF 페이지를 이미지로 렌더링한 뒤 pytesseract OCR을 수행한다."""
    try:
        import pytesseract
        from PIL import Image

        mat = mupdf_page.get_pixmap(dpi=300, alpha=False)
        img = Image.frombytes("RGB", (mat.width, mat.height), mat.samples)
        return pytesseract.image_to_string(img, lang="kor+eng")

    except Exception as e:
        logger.warning("[parser] OCR 실패: %s", e)
        return ""


def _detect_section(mupdf_page, plain_text: str) -> str | None:
    """
    페이지 첫 번째 텍스트 블록에서 섹션 제목을 감지한다.
    휴리스틱: 짧고, 굵은 폰트이거나, 큰 글자이거나, 번호로 시작하는 줄.
    """
    try:
        blocks = mupdf_page.get_text("dict").get("blocks", [])

        for block in blocks:
            if block.get("type") != 0:
                continue

            for line in block.get("lines", []):
                for span in line.get("spans", []):
                    text = span.get("text", "").strip()
                    flags = span.get("flags", 0)
                    size = span.get("size", 0)

                    if not text or len(text) > 60:
                        continue

                    is_bold = bool(flags & 2**4)
                    is_large = size >= 13
                    is_numbered = bool(re.match(r"^(\d+[.)]|[가-힣]\.|제\d+)", text))

                    if is_bold or is_large or is_numbered:
                        return text

    except Exception:
        pass

    first_line = plain_text.strip().split("\n")[0].strip() if plain_text.strip() else ""
    if first_line and len(first_line) <= 60:
        if re.match(r"^(\d+[.)]|[가-힣]\.|제\d+)", first_line):
            return first_line

    return None


# ── 비-PDF 파서 ───────────────────────────────────────────────────────────────

def _parse_docx(contents: bytes) -> list[PageResult]:
    from docx import Document

    doc = Document(io.BytesIO(contents))
    parts: list[str] = []

    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if text:
            parts.append(text)

    has_table = False
    for table in doc.tables:
        has_table = True
        rows: list[list[str]] = []

        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            rows.append(cells)

        md = _table_to_markdown(rows)
        if md:
            parts.append(md)

    return [
        PageResult(
            content="\n".join(parts),
            metadata={
                "page_number": None,
                "has_table": has_table,
                "section": None,
                "ocr_applied": False,
            },
        )
    ]


def _parse_txt(contents: bytes) -> list[PageResult]:
    text = _decode_text(contents)

    return [
        PageResult(
            content=text,
            metadata={
                "page_number": None,
                "has_table": False,
                "section": None,
                "ocr_applied": False,
            },
        )
    ]


def _parse_xlsx(contents: bytes) -> list[PageResult]:
    import pandas as pd

    xl = pd.ExcelFile(io.BytesIO(contents))
    pages: list[PageResult] = []

    for sheet_name in xl.sheet_names:
        df = xl.parse(sheet_name)
        rows = [f"[시트: {sheet_name}]"]

        for _, row in df.iterrows():
            row_parts = [
                f"{col}: {val}"
                for col, val in row.items()
                if pd.notna(val) and str(val).strip()
            ]
            if row_parts:
                rows.append(" / ".join(row_parts))

        content = "\n".join(rows)
        if content.strip():
            pages.append(
                PageResult(
                    content=content,
                    metadata={
                        "page_number": None,
                        "has_table": True,
                        "section": sheet_name,
                        "ocr_applied": False,
                    },
                )
            )

    return pages if pages else [PageResult(content="", metadata={})]


def _parse_csv(contents: bytes) -> list[PageResult]:
    import pandas as pd

    df = None
    last_error: Exception | None = None

    for encoding in ("utf-8-sig", "utf-8", "cp949", "euc-kr"):
        try:
            df = pd.read_csv(io.BytesIO(contents), encoding=encoding)
            break
        except UnicodeDecodeError as e:
            last_error = e
        except Exception as e:
            last_error = e

    if df is None:
        try:
            text = contents.decode("utf-8", errors="ignore")
            df = pd.read_csv(io.StringIO(text))
        except Exception as e:
            raise ValueError(f"CSV 파일을 읽을 수 없습니다: {e}") from last_error or e

    rows: list[str] = []
    for _, row in df.iterrows():
        row_parts = [
            f"{col}: {val}"
            for col, val in row.items()
            if pd.notna(val) and str(val).strip()
        ]
        if row_parts:
            rows.append(" / ".join(row_parts))

    return [
        PageResult(
            content="\n".join(rows),
            metadata={
                "page_number": None,
                "has_table": True,
                "section": None,
                "ocr_applied": False,
            },
        )
    ]


def _parse_pptx(contents: bytes) -> list[PageResult]:
    """
    PPTX 슬라이드별 텍스트 추출.

    처리 항목:
      - 일반 텍스트박스 / placeholder
      - GroupShape 재귀 탐색
      - 슬라이드 내 표(Table) → Markdown 변환
      - 제목 placeholder(idx=0,1) → section 메타데이터
    """
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    prs = Presentation(io.BytesIO(contents))
    pages: list[PageResult] = []

    def _extract_shape_texts(
        shapes,
        texts: list[str],
        tables_md: list[str],
        section_ref: list[str | None],
    ) -> None:
        for shape in shapes:
            if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                _extract_shape_texts(shape.shapes, texts, tables_md, section_ref)
                continue

            if getattr(shape, "has_table", False):
                rows: list[list[str]] = []
                for row in shape.table.rows:
                    cells = [cell.text.strip() for cell in row.cells]
                    rows.append(cells)

                md = _table_to_markdown(rows)
                if md:
                    tables_md.append(md)
                continue

            text = getattr(shape, "text", "").strip()
            if not text:
                continue

            if section_ref[0] is None and _is_title_placeholder(shape):
                section_ref[0] = text

            texts.append(text)

    for i, slide in enumerate(prs.slides, 1):
        texts: list[str] = []
        tables_md: list[str] = []
        section_ref: list[str | None] = [None]

        _extract_shape_texts(slide.shapes, texts, tables_md, section_ref)

        parts: list[str] = []
        if texts:
            parts.append("\n".join(texts))
        parts.extend(tables_md)

        content = f"[슬라이드 {i}]\n" + "\n\n".join(parts) if parts else ""
        if content.strip():
            pages.append(
                PageResult(
                    content=content,
                    metadata={
                        "page_number": i,
                        "has_table": len(tables_md) > 0,
                        "section": section_ref[0],
                        "ocr_applied": False,
                    },
                )
            )

    return pages if pages else [PageResult(content="", metadata={})]


def _is_title_placeholder(shape) -> bool:
    """python-pptx에서 placeholder가 아닌 shape 접근 예외를 피하면서 제목 여부를 확인한다."""
    try:
        if not getattr(shape, "is_placeholder", False):
            return False

        placeholder_format = shape.placeholder_format
        return placeholder_format.idx in (0, 1)

    except Exception:
        return False


def _decode_text(contents: bytes) -> str:
    """일반 텍스트 계열 파일을 여러 인코딩 후보로 디코딩한다."""
    for encoding in ("utf-8-sig", "utf-8", "cp949", "euc-kr"):
        try:
            return contents.decode(encoding)
        except UnicodeDecodeError:
            continue

    return contents.decode("utf-8", errors="replace")
